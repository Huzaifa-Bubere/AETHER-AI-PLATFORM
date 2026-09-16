import io
import os
from pathlib import Path
import sys
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from fastapi.testclient import TestClient
from docx import Document
import main
from services.lazy_service import LazyService
from services.resume_parser import ResumeParserService


class ServiceIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(main.app)
        self.env = patch.dict(os.environ, {'PYTHON_AI_SERVER_API_KEY': 'integration-test-only'})
        self.env.start()
        self.headers = {'Authorization': 'Bearer integration-test-only'}

    def tearDown(self):
        self.client.close()
        self.env.stop()

    def test_health_does_not_import_optional_ml(self):
        response = self.client.get('/health')
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('tensorflow', sys.modules)
        self.assertNotIn('mediapipe', sys.modules)
        self.assertNotIn('librosa', sys.modules)

    def test_missing_optional_dependency_isolated_from_resume_and_health(self):
        # Real failed import, not a substitute analysis implementation.
        with patch.object(main, 'audio_service', LazyService('ather_missing_optional_module', 'Unavailable')):
            response = self.client.post('/api/audio/analyze', headers=self.headers,
                                        json={'audio_data': 'AAAA', 'duration': 1, 'sample_rate': 44100})
            self.assertEqual(response.status_code, 503)
            health = self.client.get('/health')
            self.assertEqual(health.status_code, 200)
            self.assertEqual(health.json()['status'], 'degraded')
            self.assertEqual(self.parse_document().status_code, 200)

    def parse_document(self):
        document = Document()
        document.add_paragraph('Test Candidate')
        document.add_paragraph('Built JavaScript services with Python and SQL.')
        document.add_paragraph('Education\nBachelor of Science 2024')
        buffer = io.BytesIO()
        document.save(buffer)
        return self.client.post('/api/resume/parse', headers=self.headers,
                                files={'resume_file': ('test.docx', buffer.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})

    def test_real_docx_parse_extracts_only_present_skills(self):
        response = self.parse_document()
        self.assertEqual(response.status_code, 200)
        data = response.json()['data']
        self.assertIn('JavaScript services', data['raw_text'])
        skills = [skill for group in data['skills'] for skill in group['skills']]
        self.assertIn('javascript', skills)
        self.assertIn('python', skills)
        self.assertNotIn('java', skills)
        self.assertNotIn('r', skills)
        self.assertEqual(data['education'][0]['year'], '2024')

    def test_parser_requires_service_authentication(self):
        response = self.client.post('/api/resume/parse', files={'resume_file': ('test.txt', b'Test content')})
        self.assertIn(response.status_code, [401, 403])

    def test_skill_boundaries_keep_punctuation_languages(self):
        groups = ResumeParserService()._extract_skills('Built C++ and C# software using JavaScript; organization.')
        skills = [skill for group in groups for skill in group['skills']]
        self.assertIn('c++', skills)
        self.assertIn('c#', skills)
        self.assertNotIn('r', skills)
        self.assertNotIn('java', skills)


if __name__ == '__main__':
    unittest.main()
